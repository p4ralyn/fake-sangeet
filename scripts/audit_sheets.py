#!/usr/bin/env python
"""Audit the ticket workbooks and flag rows a human must fix.

Reports counts only -- never prints attendee data. Adds a 'NEEDS REVIEW' sheet
to the Phase 1 workbook listing rows whose buyer cannot be identified at all,
so they can be corrected at source and ingestion re-run.

Requires openpyxl only for writing that sheet; the normaliser itself is stdlib.
"""
import shutil
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from sheet_normalise import (  # noqa: E402
    Identity, TicketRow, normalise_identity, read_rows, recover_number,
    sheet_targets, split_packed,
)

SRC = Path("/Users/pragunkathuria/Downloads/ticketingestion")
EARLY = SRC / "FAKE SANGEET EARLY BIRD 2026-27.xlsx"
PHASE1 = SRC / "FAKE SANGEET_ PHASE 1 (Responses).xlsx"


def parse_early_bird() -> list[TicketRow]:
    out: list[TicketRow] = []
    for sheet, target in sheet_targets(str(EARLY)):
        if sheet not in ("INFERNO", "SNUPHORIA"):
            continue
        rows = read_rows(str(EARLY), target)
        hdr = rows[0]
        ix = {h: i for i, h in enumerate(hdr) if h}
        for n, r in enumerate(rows[1:], start=2):
            if not any((c or "").strip() for c in r):
                continue
            def cell(name):
                i = ix.get(name)
                return (r[i] if i is not None and i < len(r) else "").strip()
            tt = cell("TICKET").lower()
            is_couple = tt.startswith("duo") or tt.startswith("couple")
            ids = split_packed(cell("SNUID"))
            problems = []
            buyer = normalise_identity(ids[0]) if ids else Identity(None, None, "UNUSABLE", cell("SNUID"))
            partner = None
            mode = "SELF_CLAIM"
            if is_couple:
                if len(ids) >= 2:
                    partner = normalise_identity(ids[1])
                else:
                    # SNUPHORIA records only the buyer: nobody to send a claim
                    # link to, so both passes stay on the buyer's device.
                    mode = "PROXY"
                    problems.append("couple with no partner NetID -> PROXY")
            out.append(TicketRow(
                source=f"EARLY_BIRD/{sheet}", source_row=n,
                ticket_type="COUPLE" if is_couple else "SINGLE",
                seats=2 if is_couple else 1,
                buyer=buyer, partner=partner, partner_mode=mode,
                phone=recover_number(cell("PHONE NO.")), name=cell("NAME"),
                problems=problems,
            ))
    return out


def parse_phase1() -> list[TicketRow]:
    target = sheet_targets(str(PHASE1))[0][1]
    rows = read_rows(str(PHASE1), target)
    hdr = rows[0]
    def col(frag):
        return next(i for i, h in enumerate(hdr) if frag.lower() in h.lower())
    c_name, c_snu, c_tick = col("Name"), col("SNU Email"), col("which ticket")
    c_psnu, c_phone = col("Partner's SNU Email"), col("Phone Number")
    out = []
    for n, r in enumerate(rows[1:], start=2):
        if not any((c or "").strip() for c in r):
            continue
        def cell(i):
            return (r[i] if i < len(r) else "").strip()
        is_couple = "couple" in cell(c_tick).lower()
        buyer = normalise_identity(cell(c_snu))
        partner = normalise_identity(cell(c_psnu)) if is_couple else None
        mode = "SELF_CLAIM"
        problems = []
        if is_couple and (partner is None or not partner.usable):
            mode = "PROXY"
            problems.append("couple with no usable partner email -> PROXY")
        if buyer.status.startswith("SALVAGED"):
            problems.append(f"buyer email {buyer.status}")
        if partner and partner.status.startswith("SALVAGED"):
            problems.append(f"partner email {partner.status}")
        out.append(TicketRow(
            source="PHASE1", source_row=n,
            ticket_type="COUPLE" if is_couple else "SINGLE",
            seats=2 if is_couple else 1,
            buyer=buyer, partner=partner, partner_mode=mode,
            phone=recover_number(cell(c_phone)), name=cell(c_name),
            problems=problems,
        ))
    return out


def main() -> int:
    rows = parse_early_bird() + parse_phase1()
    ok = [r for r in rows if r.ingestable]
    bad = [r for r in rows if not r.ingestable]

    print(f"tickets parsed      : {len(rows)}")
    print(f"  ingestable        : {len(ok)}")
    print(f"  need human fix    : {len(bad)}")
    print(f"seats (ingestable)  : {sum(r.seats for r in ok)}")
    print(f"seats (all rows)    : {sum(r.seats for r in rows)}")
    print("\nby source:", dict(Counter(r.source for r in rows)))
    print("buyer id status:", dict(Counter(r.buyer.status for r in rows)))
    print("partner mode   :", dict(Counter(r.partner_mode for r in rows if r.seats == 2)))
    print("phones recovered:", sum(1 for r in rows if r.phone))

    if not bad:
        print("\nNo unusable rows. No review sheet needed.")
        return 0

    try:
        from openpyxl import load_workbook
    except ImportError:
        print("\nopenpyxl not installed; cannot write the review sheet.")
        print("Rows needing a fix (by source row):",
              ", ".join(f"{r.source}:{r.source_row}" for r in bad))
        return 1

    backup = PHASE1.with_suffix(".backup.xlsx")
    if not backup.exists():
        shutil.copy2(PHASE1, backup)
        print(f"\nbacked up original -> {backup.name}")

    wb = load_workbook(PHASE1)
    if "NEEDS REVIEW" in wb.sheetnames:
        del wb["NEEDS REVIEW"]
    ws = wb.create_sheet("NEEDS REVIEW")
    ws.append(["Source", "Row in sheet", "Name", "Ticket",
               "What we read", "Problem", "What to do"])
    for r in bad:
        ws.append([
            r.source, r.source_row, r.name, r.ticket_type, r.buyer.raw,
            "No NetID could be read from the SNU email cell",
            f"Put the person's NetID or full @snu.edu.in address in row "
            f"{r.source_row}, then re-run ingestion.",
        ])
    for col_cells, width in zip(ws.columns, (18, 12, 26, 10, 30, 42, 60)):
        ws.column_dimensions[col_cells[0].column_letter].width = width
    wb.save(PHASE1)
    print(f"wrote 'NEEDS REVIEW' sheet with {len(bad)} rows -> {PHASE1.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
